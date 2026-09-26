"""One-pager PPTX for the Second Look pitch: concept, AI role, business value."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- palette ----
NAVY = RGBColor(0x1B, 0x23, 0x40)
NAVY_SOFT = RGBColor(0xCA, 0xDC, 0xFC)   # light ice-blue, for text on navy
CARD_BG = RGBColor(0xED, 0xF1, 0xF8)
AMBER = RGBColor(0xF2, 0xA9, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x55, 0x5F, 0x75)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def no_line(shape):
    shape.line.fill.background()


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    no_line(shape)


def set_round_radius(shape, frac):
    """0..0.5 adjustment for ROUNDED_RECTANGLE corner radius."""
    try:
        shape.adjustments[0] = frac
    except Exception:
        pass


def add_text(x, y, w, h, text, size, color, bold=False, italic=False, font=BODY_FONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_bullets(x, y, w, h, items, size, color, font=BODY_FONT, space_after=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.08
        p.space_after = Pt(space_after)
        r_bullet = p.add_run()
        r_bullet.text = "●  "
        r_bullet.font.size = Pt(size - 2)
        r_bullet.font.color.rgb = AMBER
        r_bullet.font.name = font
        r_bullet.font.bold = True
        r_text = p.add_run()
        r_text.text = item
        r_text.font.size = Pt(size)
        r_text.font.color.rgb = color
        r_text.font.name = font
    return box


# ---------------------------------------------------------------- header
header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(SW), Inches(1.35))
solid(header, NAVY)
header.shadow.inherit = False

add_text(0.5, 0.20, 8.6, 0.62, "Second Look", 34, WHITE, bold=True, font=HEAD_FONT)
add_text(0.5, 0.82, 8.6, 0.42,
         "Rescuing qualified candidates a keyword ATS wrongly rejected",
         15, NAVY_SOFT, italic=True, font=BODY_FONT)

badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.15), Inches(0.36), Inches(2.68), Inches(0.62))
solid(badge, AMBER)
set_round_radius(badge, 0.5)
badge.shadow.inherit = False
btf = badge.text_frame
btf.word_wrap = True
btf.vertical_anchor = MSO_ANCHOR.MIDDLE
btf.margin_left = Inches(0.05)
btf.margin_right = Inches(0.05)
btf.margin_top = 0
btf.margin_bottom = 0
p1 = btf.paragraphs[0]
p1.alignment = PP_ALIGN.CENTER
r1 = p1.add_run()
r1.text = "PITCHFEST 2026"
r1.font.size = Pt(11)
r1.font.bold = True
r1.font.name = BODY_FONT
r1.font.color.rgb = NAVY
p2 = btf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(0)
r2 = p2.add_run()
r2.text = "ENTERPRISE AI TRACK"
r2.font.size = Pt(9)
r2.font.bold = True
r2.font.name = BODY_FONT
r2.font.color.rgb = NAVY

# ---------------------------------------------------------------- three columns
COL_Y = 1.65
COL_H = 4.85
COL_W = 3.9
GAP = 0.3
COL_X = [0.5, 0.5 + COL_W + GAP, 0.5 + 2 * (COL_W + GAP)]

titles = ["The problem", "How AI does it", "Business value"]


def card(cx):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(COL_Y), Inches(COL_W), Inches(COL_H))
    solid(rect, CARD_BG)
    set_round_radius(rect, 0.04)
    rect.shadow.inherit = False
    return rect


def badge_circle(cx, number):
    d = 0.6
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.3), Inches(COL_Y + 0.3), Inches(d), Inches(d))
    solid(oval, NAVY)
    oval.shadow.inherit = False
    tf = oval.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(number)
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = BODY_FONT


for i, cx in enumerate(COL_X):
    card(cx)
    badge_circle(cx, i + 1)
    add_text(cx + 1.05, COL_Y + 0.30, COL_W - 1.35, 0.6, titles[i], 19, NAVY, bold=True,
              font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)

# --- column 1: the problem
add_bullets(
    COL_X[0] + 0.3, COL_Y + 1.15, COL_W - 0.6, COL_H - 1.45,
    [
        "Keyword ATS filters reject real skill over how it's phrased",
        "Enterprises lose candidates they already paid to attract",
        "Often tracks age/gap bias — see: Mobley v. Workday",
        "Existing tools mostly rank new applicants — few re-audit past rejections",
    ],
    14, NAVY,
)

# --- column 2: how AI does it
add_bullets(
    COL_X[1] + 0.3, COL_Y + 1.15, COL_W - 0.6, COL_H - 1.45,
    [
        "Extracts “keywords” (what ATS scans) vs. “equivalents” it misses",
        "Per CV: verbatim evidence quote, bias flags, 3 interview questions",
        "Code — not the model — verifies quotes & computes the score",
        "Live-proven: paste a new JD and the rubric & shortlist change",
        "Reproducible: 112/112 verdicts identical across independent runs",
    ],
    13.5, NAVY,
)

# --- column 3: business value (stats row + bullets)
STAT_Y = COL_Y + 1.15
stat_w = (COL_W - 0.6 - 2 * 0.15) / 3
stats = [("71%", "hard-rescue\nrecall"), ("0", "false\nrescues"), ("99%", "quotes\nverified")]
for j, (num, cap) in enumerate(stats):
    sx = COL_X[2] + 0.3 + j * (stat_w + 0.15)
    add_text(sx, STAT_Y, stat_w, 0.55, num, 26, AMBER, bold=True, font=HEAD_FONT,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    add_text(sx, STAT_Y + 0.55, stat_w, 0.5, cap, 9.5, MUTED, font=BODY_FONT,
              align=PP_ALIGN.CENTER)

add_bullets(
    COL_X[2] + 0.3, STAT_Y + 1.25, COL_W - 0.6, COL_H - (1.25 + STAT_Y - COL_Y) - 0.3,
    [
        "Recovers qualified people — without lowering the bar",
        "Audit trail: evidence, rejection-pattern report, CSV export",
        "Always human-in-the-loop — recommends, never decides",
    ],
    14, NAVY,
)

# ---------------------------------------------------------------- footer
add_text(
    0.5, 6.85, 12.333, 0.4,
    "Stack: Python · Streamlit · Google Gemini (structured output) · grounding & scoring computed in code · no GPU needed",
    10.5, MUTED, italic=True, font=BODY_FONT, align=PP_ALIGN.CENTER,
)

prs.save("Second_Look_One_Pager.pptx")
print("wrote Second_Look_One_Pager.pptx")
