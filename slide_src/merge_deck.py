"""Merge the standalone one-pager decks into a single submission deck.

Pure structural copy: each source slide's shape tree (autoshapes/textboxes, no
images/charts in either source) is deep-copied onto a fresh blank slide, and its
notes text is copied across. Content is not edited.
"""
import copy
import sys

from pptx import Presentation
from pptx.util import Emu

OUTPUT = "BPM_2026_Day1_PitchDemo_VibeTribe.pptx"
SOURCES = ["Relook_One_Pager.pptx", "Baton_One_Pager.pptx"]

first = Presentation(SOURCES[0])
merged = Presentation()
merged.slide_width = first.slide_width
merged.slide_height = first.slide_height
blank_layout = merged.slide_layouts[6]

total = 0
for src_path in SOURCES:
    src = Presentation(src_path)
    if (src.slide_width, src.slide_height) != (merged.slide_width, merged.slide_height):
        print(f"WARNING: {src_path} is {Emu(src.slide_width).inches}x{Emu(src.slide_height).inches}in, "
              f"expected {Emu(merged.slide_width).inches}x{Emu(merged.slide_height).inches}in", file=sys.stderr)
    for src_slide in src.slides:
        new_slide = merged.slides.add_slide(blank_layout)
        for shape in src_slide.shapes:
            new_slide.shapes._spTree.append(copy.deepcopy(shape._element))
        if src_slide.has_notes_slide and src_slide.notes_slide.notes_text_frame.text.strip():
            new_slide.notes_slide.notes_text_frame.text = src_slide.notes_slide.notes_text_frame.text
        total += 1
        print(f"  + slide {total}: from {src_path} ({len(src_slide.shapes)} shapes)")

merged.save(OUTPUT)
print(f"wrote {OUTPUT} ({total} slides)")
