"""One-pager PPTX for the Baton pitch: concept, workflow, AI role, business value.

Single slide (previously split as overview + workflow-detail across two slides;
consolidated here so the submission deck carries one Baton slide, not two)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- palette ----
NAVY = RGBColor(0x1B, 0x23, 0x40)
NAVY_SOFT = RGBColor(0xCA, 0xDC, 0xFC)   # light ice-blue, for text on navy
CARD_BG = RGBColor(0xED, 0xF1, 0xF8)
AMBER = RGBColor(0xF2, 0xA9, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x55, 0x5F, 0x75)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def no_line(shape):
    shape.line.fill.background()


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    no_line(shape)


def set_round_radius(shape, frac):
    """0..0.5 adjustment for ROUNDED_RECTANGLE corner radius."""
    try:
        shape.adjustments[0] = frac
    except Exception:
        pass


def add_text(x, y, w, h, text, size, color, bold=False, italic=False, font=BODY_FONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_bullets(x, y, w, h, items, size, color, font=BODY_FONT, space_after=5, line_spacing=1.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r_bullet = p.add_run()
        r_bullet.text = "●  "
        r_bullet.font.size = Pt(size - 2)
        r_bullet.font.color.rgb = AMBER
        r_bullet.font.name = font
        r_bullet.font.bold = True
        r_text = p.add_run()
        r_text.text = item
        r_text.font.size = Pt(size)
        r_text.font.color.rgb = color
        r_text.font.name = font
    return box


# ---------------------------------------------------------------- header
header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(SW), Inches(0.95))
solid(header, NAVY)
header.shadow.inherit = False

add_text(0.5, 0.13, 8.6, 0.5, "Baton", 28, WHITE, bold=True, font=HEAD_FONT)
add_text(0.5, 0.62, 8.6, 0.3, "Knowledge transfer you can prove is complete", 13, NAVY_SOFT,
         italic=True, font=BODY_FONT)

badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.35), Inches(0.20), Inches(2.48), Inches(0.56))
solid(badge, AMBER)
set_round_radius(badge, 0.5)
badge.shadow.inherit = False
btf = badge.text_frame
btf.word_wrap = True
btf.vertical_anchor = MSO_ANCHOR.MIDDLE
btf.margin_left = Inches(0.05)
btf.margin_right = Inches(0.05)
btf.margin_top = 0
btf.margin_bottom = 0
p1 = btf.paragraphs[0]
p1.alignment = PP_ALIGN.CENTER
r1 = p1.add_run()
r1.text = "PITCHFEST 2026"
r1.font.size = Pt(10)
r1.font.bold = True
r1.font.name = BODY_FONT
r1.font.color.rgb = NAVY
p2 = btf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(0)
r2 = p2.add_run()
r2.text = "ENTERPRISE AI TRACK"
r2.font.size = Pt(8)
r2.font.bold = True
r2.font.name = BODY_FONT
r2.font.color.rgb = NAVY

# ---------------------------------------------------------------- three columns
COL_Y = 1.10
COL_H = 4.05
COL_W = 3.9
GAP = 0.3
COL_X = [0.5, 0.5 + COL_W + GAP, 0.5 + 2 * (COL_W + GAP)]

titles = ["The problem", "How AI does it", "Business value"]


def card(cx):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(COL_Y), Inches(COL_W), Inches(COL_H))
    solid(rect, CARD_BG)
    set_round_radius(rect, 0.04)
    rect.shadow.inherit = False
    return rect


def badge_circle(cx, number):
    d = 0.52
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.26), Inches(COL_Y + 0.26), Inches(d), Inches(d))
    solid(oval, NAVY)
    oval.shadow.inherit = False
    tf = oval.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(number)
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = BODY_FONT


for i, cx in enumerate(COL_X):
    card(cx)
    badge_circle(cx, i + 1)
    add_text(cx + 0.92, COL_Y + 0.26, COL_W - 1.2, 0.52, titles[i], 16.5, NAVY, bold=True,
              font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)

# --- column 1: the problem
add_bullets(
    COL_X[0] + 0.28, COL_Y + 0.98, COL_W - 0.56, COL_H - 1.2,
    [
        "KT is a meeting and a checklist: nobody tests whether the giver and taker actually agree",
        "Gaps surface months later as incidents and slow ramp-up",
        "Project managers can't see KT status, blockers, or who is holding it up",
    ],
    13, NAVY,
)

# --- column 2: how AI does it (workflow + AI/code/human roles, condensed)
add_bullets(
    COL_X[1] + 0.28, COL_Y + 0.98, COL_W - 0.56, COL_H - 1.2,
    [
        "AI asks giver and taker separately; round 2 follows up on weak angles only",
        "Code verifies quotes and scores readiness against the manager's target",
        "AI recommends and reviews action items; code decides close / redress / escalate",
        "Full audit trail: every question, answer, action and decision is logged",
    ],
    12.5, NAVY,
)

# --- column 3: business value (stats row + bullets)
STAT_Y = COL_Y + 0.98
stat_w = (COL_W - 0.56 - 2 * 0.12) / 3
stats = [("4/4", "planted gaps\ndetected"), ("97%", "final readiness,\nup from 57%"), ("0/8", "false alarms on\nsound angles")]
for j, (num, cap) in enumerate(stats):
    sx = COL_X[2] + 0.28 + j * (stat_w + 0.12)
    add_text(sx, STAT_Y, stat_w, 0.46, num, 21, AMBER, bold=True, font=HEAD_FONT,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    add_text(sx, STAT_Y + 0.46, stat_w, 0.42, cap, 8.5, MUTED, font=BODY_FONT,
              align=PP_ALIGN.CENTER, line_spacing=1.0)

add_bullets(
    COL_X[2] + 0.28, STAT_Y + 1.02, COL_W - 0.56, COL_H - (1.02 + STAT_Y - COL_Y) - 0.15,
    [
        "KT closes on evidence, not on a signed checklist",
        "Manager tracker: stage, day n of N, deadline, who each item waits on",
        "Turns a KT into an auditable, escalation-ready workflow",
    ],
    13, NAVY,
)

# ---------------------------------------------------------------- 6-step workflow strip
flow = [
    ("1 Brief", "PM enters project, roles, stack, dev/test env, duration"),
    ("2 Q&A", "Giver & taker answer independently across 12 KT angles"),
    ("3 Assess", "AI compares both sides; quotes verified in code"),
    ("4 Actions", "Gaps become owned action items, with proof or justification"),
    ("5 Review", "AI reviews responses; weak justifications rejected"),
    ("6 Decide", "Code: Close, one redress round, or Escalate"),
]
FN = len(flow)
FG = 0.13
FW = (12.333 - (FN - 1) * FG) / FN
FY, FH = COL_Y + COL_H + 0.13, 1.00
for i, (t, d) in enumerate(flow):
    fx = 0.5 + i * (FW + FG)
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(fx), Inches(FY), Inches(FW), Inches(FH))
    solid(r, CARD_BG)
    set_round_radius(r, 0.06)
    r.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(fx), Inches(FY), Inches(FW), Inches(0.32))
    solid(bar, NAVY if i < FN - 1 else AMBER)
    bar.shadow.inherit = False
    add_text(fx + 0.1, FY + 0.02, FW - 0.16, 0.28, t, 11.5, WHITE if i < FN - 1 else NAVY,
             bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(fx + 0.1, FY + 0.38, FW - 0.2, FH - 0.42, d, 9, NAVY, line_spacing=1.05)

add_text(0.5, FY + FH + 0.03, 12.333, 0.22,
         "Redress loop: a rejected item returns to its owner once, with the reviewer's note. Still unresolved: escalated to the manager.",
         9.5, MUTED, italic=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- footer
add_text(
    0.5, 6.58, 12.333, 0.3,
    "Stack: Python · Streamlit · Google Gemini (structured output), Groq fallback on outage/quota · no GPU needed",
    10, MUTED, italic=True, font=BODY_FONT, align=PP_ALIGN.CENTER,
)
add_text(
    0.5, 6.87, 12.333, 0.3,
    "Eval: simulated giver/taker personas, 4 planted weak angles, one synthetic project, one run (69s end to end). Target is a stopping rule, not proof of completeness.",
    9, MUTED, italic=True, font=BODY_FONT, align=PP_ALIGN.CENTER,
)

prs.save("Baton_One_Pager.pptx")
print("wrote Baton_One_Pager.pptx (1 slide)")
