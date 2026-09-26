"""F2: turn uploaded CV files (PDF / DOCX / TXT / MD) into Candidate records.

Only text-based files are supported. A scanned/image-only PDF has no extractable
text layer and is reported back as a warning rather than silently scored as empty.
"""
from __future__ import annotations

import io
import re
from pathlib import PurePath

from app.schemas import Candidate

SUPPORTED_EXTENSIONS = ("pdf", "docx", "txt", "md")
MIN_CHARS = 200  # anything shorter is almost certainly a failed extraction, not a CV


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Many CVs put skills/experience in tables; include cell text too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        text = _pdf_text(data)
    elif ext == "docx":
        text = _docx_text(data)
    elif ext in ("txt", "md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: .{ext} (use {', '.join(SUPPORTED_EXTENSIONS)})")
    return re.sub(r"[ \t]+\n", "\n", text).strip()


def candidates_from_uploads(files: list[tuple[str, bytes]]) -> tuple[list[Candidate], list[str]]:
    """files = [(filename, raw_bytes), ...]. Returns (candidates, warnings)."""
    candidates: list[Candidate] = []
    warnings: list[str] = []
    for i, (filename, data) in enumerate(files, start=1):
        try:
            text = extract_text(filename, data)
        except Exception as e:  # noqa: BLE001 - surface any parser failure per-file
            warnings.append(f"{filename}: could not be read ({type(e).__name__}: {e})")
            continue
        if len(text) < MIN_CHARS:
            warnings.append(
                f"{filename}: only {len(text)} characters of text found - skipped "
                "(a scanned/image-only PDF has no text layer to read)"
            )
            continue
        stem = PurePath(filename).stem
        name = re.sub(r"[_\-]+", " ", stem).strip().title() or f"Upload {i}"
        candidates.append(Candidate(id=f"U{len(candidates) + 1:02d}", name=name, text=text))
    return candidates, warnings
